"""Generate real, validated PDF and DOCX from the saved section order."""
import io
from html import escape
from html.parser import HTMLParser
from pathlib import Path
import threading
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
import pymupdf
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from .editor_content import DraftContent, clean_html
from .templates import TEMPLATES


class Blocks(HTMLParser):
    def __init__(self):
        super().__init__()
        self.blocks = []
        self.runs = []
        self.marks = []
        self.href = None
        self.list_depth = 0
        self.is_bullet = False

    def flush(self):
        if self.runs:
            self.blocks.append({'runs': self.runs, 'bullet': self.is_bullet})
        self.runs = []

    def handle_starttag(self, tag, attrs):
        if tag == 'li':
            self.flush()
            self.list_depth += 1
            self.is_bullet = True
        elif tag in ('p', 'h2', 'h3'):
            self.flush()
            self.is_bullet = self.list_depth > 0
        if tag in ('strong', 'em', 'u'):
            self.marks.append(tag)
        if tag == 'a':
            self.href = dict(attrs).get('href')
        if tag == 'br':
            self.handle_data('\n')

    def handle_endtag(self, tag):
        if tag in ('p', 'h2', 'h3', 'li'):
            self.flush()
        if tag == 'li':
            self.list_depth = max(0, self.list_depth - 1)
            self.is_bullet = self.list_depth > 0
        if tag in self.marks:
            self.marks.remove(tag)
        if tag == 'a':
            self.href = None

    def handle_data(self, value):
        if value:
            self.runs.append({'text': value, 'bold': 'strong' in self.marks, 'italic': 'em' in self.marks, 'underline': 'u' in self.marks, 'href': self.href})


def blocks(value):
    parser = Blocks()
    parser.feed(clean_html(value))
    parser.flush()
    return parser.blocks


def hyperlink(paragraph, run):
    rel = paragraph.part.relate_to(run['href'], RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), rel)
    node = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    for flag, name in [('bold', 'w:b'), ('italic', 'w:i'), ('underline', 'w:u')]:
        if run[flag]:
            element = OxmlElement(name)
            if flag == 'underline': element.set(qn('w:val'), 'single')
            props.append(element)
    node.append(props)
    value = OxmlElement('w:t')
    value.set(qn('xml:space'), 'preserve')
    value.text = run['text']
    node.append(value)
    link.append(node)
    paragraph._p.append(link)


def docx_bytes(content: DraftContent):
    template = TEMPLATES[content.template]
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Pt(595.28), Pt(841.89)
    section.top_margin = section.bottom_margin = Pt(template['margin'])
    section.left_margin = section.right_margin = Pt(template['margin'])
    style = document.styles['Normal']
    style.font.name = 'Times New Roman' if template['font'] == 'serif' else 'Calibri'
    style.font.size = Pt(template['body_size'])
    style.paragraph_format.space_after = Pt(template['spacing'])
    style.paragraph_format.line_spacing = Pt(template['leading'])
    heading_style = document.styles['Heading 1']
    heading_style.font.name = style.font.name
    heading_style.font.size = Pt(template['heading_size'])
    heading_style.font.color.rgb = RGBColor.from_string(template['accent'].lstrip('#'))
    heading_style.paragraph_format.keep_with_next = True
    for item in content.sections:
        if item.title.lower() != 'contact information':
            heading = document.add_heading(item.title, level=1)
            heading.paragraph_format.keep_with_next = True
        for block in blocks(item.html):
            p = document.add_paragraph(style='List Bullet' if block['bullet'] else None)
            for value in block['runs']:
                if value['href']:
                    hyperlink(p, value)
                else:
                    run = p.add_run(value['text'])
                    run.bold, run.italic, run.underline = value['bold'], value['italic'], value['underline']
    buffer = io.BytesIO()
    document.save(buffer)
    data = buffer.getvalue()
    # Reopen the actual generated file, not just a success response.
    reopened = Document(io.BytesIO(data))
    if not reopened.paragraphs:
        raise ValueError('The exported DOCX is empty.')
    return data


_font_lock = threading.Lock()


def pdf_font(family='sans'):
    base = 'FormiqoSerif' if family == 'serif' else 'Formiqo'
    with _font_lock:
        if base in pdfmetrics.getRegisteredFontNames():
            return base
        families = [
            (Path('/usr/share/fonts/truetype/dejavu'), ['DejaVuSans.ttf', 'DejaVuSans-Bold.ttf', 'DejaVuSans-Oblique.ttf', 'DejaVuSans-BoldOblique.ttf']),
            (Path('/System/Library/Fonts/Supplemental'), ['Arial.ttf', 'Arial Bold.ttf', 'Arial Italic.ttf', 'Arial Bold Italic.ttf']),
        ]
        if family == 'serif':
            families = [
                (Path('/usr/share/fonts/truetype/dejavu'), ['DejaVuSerif.ttf', 'DejaVuSerif-Bold.ttf', 'DejaVuSerif-Italic.ttf', 'DejaVuSerif-BoldItalic.ttf']),
                (Path('/System/Library/Fonts/Supplemental'), ['Times New Roman.ttf', 'Times New Roman Bold.ttf', 'Times New Roman Italic.ttf', 'Times New Roman Bold Italic.ttf']),
            ]
        for directory, names in families:
            if all((directory / name).exists() for name in names):
                for suffix, name in zip(['', '-Bold', '-Italic', '-BoldItalic'], names):
                    pdfmetrics.registerFont(TTFont(base + suffix, str(directory / name)))
                pdfmetrics.registerFontFamily(base, normal=base, bold=base+'-Bold', italic=base+'-Italic', boldItalic=base+'-BoldItalic')
                return base
    return 'Times-Roman' if family == 'serif' else 'Helvetica'


def pdf_bytes(content: DraftContent):
    template = TEMPLATES[content.template]
    buffer = io.BytesIO()
    font = pdf_font(template['font'])
    styles = getSampleStyleSheet()
    body = ParagraphStyle('Body', parent=styles['Normal'], fontName=font, fontSize=template['body_size'], leading=template['leading'], spaceAfter=template['spacing'], splitLongWords=True)
    bold_font = font+'-Bold' if font.startswith('Formiqo') else 'Times-Bold' if font == 'Times-Roman' else 'Helvetica-Bold'
    heading = ParagraphStyle('Section', parent=body, fontName=bold_font, fontSize=template['heading_size'], textColor=HexColor(template['accent']), leading=template['heading_size'] + 5, spaceBefore=12, spaceAfter=template['spacing'] + 1, keepWithNext=True)
    flow = []
    for section in content.sections:
        if section.title.lower() != 'contact information':
            flow.append(Paragraph(escape(section.title), heading))
        for block in blocks(section.html):
            parts = []
            for run in block['runs']:
                value = escape(run['text']).replace('\n', '<br/>')
                if run['bold']: value = '<b>' + value + '</b>'
                if run['italic']: value = '<i>' + value + '</i>'
                if run['underline']: value = '<u>' + value + '</u>'
                if run['href']: value = '<link href="' + escape(run['href'], quote=True) + '">' + value + '</link>'
                parts.append(value)
            flow.append(Paragraph(''.join(parts), body, bulletText='•' if block['bullet'] else None))
        flow.append(Spacer(1, 3))
    margin = template['margin']
    document = SimpleDocTemplate(buffer, pagesize=(595.28, 841.89), leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=margin, title='Formiqo Resume', author='Formiqo')
    document.build(flow)
    data = buffer.getvalue()
    with pymupdf.open(stream=data, filetype='pdf') as pdf:
        if len(pdf) == 0 or not any(p.get_text().strip() for p in pdf):
            raise ValueError('The exported PDF has no readable text.')
    return data
