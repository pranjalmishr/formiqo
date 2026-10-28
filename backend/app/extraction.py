import re
import shutil
import subprocess
import tempfile
from pathlib import Path
import pymupdf as fitz
from docx import Document
from .config import settings
from .pdf_layout import analyze_document


class ExtractionError(ValueError):
    pass


def is_section_heading(text: str):
    from .scoring import heading_key
    return heading_key(text) is not None


def normalize(text: str):
    text = text.replace('\r\n', '\n').replace('\r', '\n').replace('\x00', '')
    text = '\n'.join(re.sub(r'[ \t]+', ' ', line).strip() for line in text.split('\n'))
    return re.sub(r'\n{4,}', '\n\n\n', text).strip()


def ocr_page(page):
    if not shutil.which('tesseract'):
        raise ExtractionError('Scanned page detected, but Tesseract OCR is not installed. Install Tesseract or upload a selectable-text PDF, DOCX, or TXT.')
    # OCR runs without a shell and under the extraction worker timeout.
    with tempfile.TemporaryDirectory(prefix='resumeai-ocr-') as tmp:
        image = Path(tmp) / 'page.png'
        page.get_pixmap(dpi=150, alpha=False).save(image)
        result = subprocess.run(['tesseract', str(image), 'stdout', '-l', 'eng'], capture_output=True, timeout=30, check=False)
        if result.returncode:
            raise ExtractionError('OCR failed. Try a clearer scan or paste your resume into a TXT file.')
        return result.stdout.decode('utf-8', errors='replace')


def extract(path: Path, ext: str):
    warnings = []
    metadata = {'format': ext, 'page_count': None, 'ocr_pages': [], 'warnings': warnings, 'checks': [], 'limited_layout_detection': True}
    parts = []
    if ext == 'pdf':
        try:
            with fitz.open(path) as pdf:
                if pdf.needs_pass:
                    raise ExtractionError('Password-protected PDFs are not supported. Upload an unlocked copy.')
                if len(pdf) > settings.max_pages:
                    raise ExtractionError(f'Resumes are limited to {settings.max_pages} pages.')
                metadata['page_count'] = len(pdf)
                for page in pdf:
                    if page.rect.width * page.rect.height > 4_000_000:
                        raise ExtractionError('PDF page dimensions exceed the processing limit.')
                # Column-aware reading order (sidebars are read as a unit, wrapped lines are
                # rejoined, page numbers and repeated running headers are dropped). A
                # geometry failure falls back to spatial sorting rather than losing text.
                try:
                    layouts = {layout.number: layout for layout in analyze_document(pdf, is_section_heading)}
                except Exception:
                    layouts = {}
                removed = []
                for index, page in enumerate(pdf):
                    layout = layouts.get(index + 1)
                    text = layout.text() if layout else page.get_text('text', sort=True)
                    if len(text.strip()) < 30 and page.get_images():
                        metadata['ocr_pages'].append(index + 1)
                        text = ocr_page(page)
                    elif layout:
                        removed.extend(layout.removed)
                    parts.append(text)
                    blocks = [b for b in page.get_text('blocks') if len(b) > 6 and b[6] == 0]
                    if layout and layout.columns == 2:
                        warnings.append({'code': 'columns', 'page': index + 1, 'message': 'A two-column or sidebar layout was detected and read one column at a time. Some ATS parsers read across columns instead; verify the extracted order.'})
                    elif not layout:
                        left = [b for b in blocks if b[2] <= page.rect.width * 0.53]
                        right = [b for b in blocks if b[0] >= page.rect.width * 0.45]
                        if any(abs(a[1] - b[1]) < 25 and a[2] < b[0] for a in left for b in right):
                            warnings.append({'code': 'columns', 'page': index + 1, 'message': 'Separated side-by-side text blocks may indicate columns. Verify reading order in the extracted preview.'})
                    for block in page.get_text('dict')['blocks']:
                        for line in block.get('lines', []):
                            for span in line.get('spans', []):
                                if span.get('text', '').strip() and span.get('size', 12) < 8:
                                    warnings.append({'code': 'small_font', 'page': index + 1, 'message': 'Some text is smaller than 8 pt. Verify readability.'})
                                    break
                    edge_text = ' '.join(b[4] for b in blocks if b[1] < 35 or b[3] > page.rect.height - 35)
                    if re.search(r'[\w.+-]+@[\w.-]+\.[a-zA-Z]{2,}', edge_text):
                        warnings.append({'code': 'edge_contact', 'page': index + 1, 'message': 'An email appears near a page edge; check that contact details are in the main reading flow.'})
                metadata['checks'] = ['Page order preserved', 'Selectable text checked', 'Column-aware reading order applied', 'Small-font and column heuristics checked']
                if removed:
                    metadata['removed_artifacts'] = list(dict.fromkeys(removed))[:40]
                    metadata['checks'].append(f'{len(removed)} page-number/running-header line(s) excluded from text')
        except ExtractionError:
            raise
        except Exception as exc:
            raise ExtractionError('This PDF could not be read. It may be corrupted. Try re-exporting it.') from exc
    elif ext == 'docx':
        try:
            doc = Document(path)
            # Walk body children to retain paragraph/table order; never execute embedded content.
            from docx.text.paragraph import Paragraph
            from docx.table import Table
            for child in doc.element.body.iterchildren():
                if child.tag.endswith('}p'):
                    para = Paragraph(child, doc)
                    text = para.text
                    if para._p.xpath('.//w:numPr') and text.strip():
                        text = '• ' + text
                    parts.append(text)
                elif child.tag.endswith('}tbl'):
                    table = Table(child, doc)
                    parts.extend(' | '.join(c.text for c in row.cells) for row in table.rows)
                    warnings.append({'code': 'tables', 'message': 'Table content was extracted row by row. Essential information in tables may have an unusual reading order.'})
            for section in doc.sections:
                extra = '\n'.join(p.text for p in list(section.header.paragraphs) + list(section.footer.paragraphs)).strip()
                if extra:
                    parts.append(extra)
                    warnings.append({'code': 'header_footer', 'message': 'Header/footer text was appended after body content. Check that essential details are not missed.'})
            if doc.element.body.xpath('.//w:txbxContent'):
                warnings.append({'code': 'text_boxes', 'message': 'Text boxes detected; their content may not be fully extracted. Confirm the preview.'})
            metadata['checks'] = ['Body paragraph/table order preserved', 'Headers, footers and text boxes checked']
        except Exception as exc:
            raise ExtractionError('This DOCX could not be read. Try saving a new DOCX copy.') from exc
    else:
        parts = [path.read_text(encoding='utf-8-sig')]
        metadata['checks'] = ['UTF-8 plain text validated']
    text = normalize('\n\n'.join(parts))
    if not text:
        raise ExtractionError('No readable text was found. Upload a clearer copy or a TXT resume.')
    if len(text) > settings.max_text_chars:
        raise ExtractionError('Extracted text exceeds the resume text limit.')
    if metadata['ocr_pages']:
        warnings.append({'code': 'ocr', 'message': 'Image-based pages were processed with OCR. Review spelling and reading order before analysis.'})
    metadata['warnings'] = list({(w['code'], w.get('page')): w for w in warnings}.values())
    try:
        from .resume_import import profile_source
        metadata['reconstruction'] = profile_source(path,ext,text,metadata['ocr_pages'])
    except Exception:
        metadata['warnings'].append({'code':'reconstruction_limited','message':'Visual structure could not be reconstructed safely. The complete extracted text and retained original remain available for review.'})
    return {'text': text, 'extraction': metadata}
