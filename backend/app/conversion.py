"""Real conversion matrix and bounded file transformations. No remote resource loading.

PDF->DOCX is reconstruction, not a claimed pixel-identical replica. General
DOCX->PDF now requires actual Office rendering: the former generic-text fallback
is deliberately not advertised. Office uses enabled LibreOffice + bubblewrap.
"""
import io
import os
import re
import shutil
import subprocess
import zipfile
from collections import Counter
from html import escape, unescape
from pathlib import Path
from typing import Literal
from functools import lru_cache
from pydantic import BaseModel, Field
import bleach
import pymupdf
from docx import Document
from docx.shared import Pt
from PIL import Image, ImageOps, features
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from .config import settings
from .editor_content import DraftContent, Section, clean_html, plain_text, text_to_html
from .exports import pdf_bytes, docx_bytes
from .extraction import extract, ocr_page, ExtractionError

Image.MAX_IMAGE_PIXELS = 20_000_000
IMAGE_TYPES = {'jpg': 'JPEG', 'png': 'PNG', 'webp': 'WEBP'}
MIMES = {'pdf':'application/pdf','docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','txt':'text/plain','jpg':'image/jpeg','png':'image/png','webp':'image/webp','html':'text/html','zip':'application/zip','pptx':'application/vnd.openxmlformats-officedocument.presentationml.presentation','xlsx':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','doc':'application/msword'}


@lru_cache(maxsize=4)
def office_preflight(executable, sandbox):
    if not executable or not sandbox: return False
    command=[sandbox,'--unshare-all','--die-with-parent','--clearenv','--ro-bind','/usr','/usr','--proc','/proc','--dev','/dev']
    for directory in ['/lib','/lib64','/bin']:
        if Path(directory).exists(): command+=['--ro-bind',directory,directory]
    try:
        return subprocess.run(command+['--','/bin/true'],capture_output=True,timeout=5,check=False).returncode==0
    except (OSError,subprocess.TimeoutExpired): return False


def office_available():
    return bool(settings.office_enabled and os.name=='posix' and office_preflight(shutil.which('libreoffice'),shutil.which('bwrap')))


def matrix():
    result = {'pdf':['docx','txt','jpg','png'], 'docx':['txt','html'], 'txt':['docx','pdf'], 'jpg':['png','pdf'], 'png':['jpg','pdf'], 'html':['pdf','txt']}
    if features.check('webp'):
        result['webp']=['jpg','png','pdf']; result['jpg'].append('webp'); result['png'].append('webp')
    if office_available():
        result.update({'doc':['pdf'], 'pptx':['pdf'], 'xlsx':['pdf']})
        result['docx'].insert(0,'pdf')
    return result


class ConversionSettings(BaseModel):
    model_config = {'extra':'forbid'}
    dpi: Literal[72,150,200] = 150
    pages: str = Field(default='', max_length=150)
    page_size: Literal['A4','LETTER'] = 'A4'
    orientation: Literal['portrait','landscape'] = 'portrait'
    margin_mm: int = Field(default=12, ge=0, le=30)


class ConversionError(ValueError):
    def __init__(self, message, code='processing_failed'):
        super().__init__(message); self.code=code


def page_indices(value, total):
    if not value.strip(): return list(range(total))
    selected=set()
    for group in value.split(','):
        group=group.strip()
        match=re.fullmatch(r'(\d+)(?:\s*-\s*(\d+))?',group)
        if not match: raise ConversionError('Use page ranges such as 1-3,5.', 'invalid_page_range')
        start=int(match[1]); end=int(match[2] or match[1])
        if not 1 <= start <= end <= total: raise ConversionError('The selected page range is outside this PDF.', 'invalid_page_range')
        selected.update(range(start-1,end))
    return sorted(selected)


def safe_text(path):
    try: text=path.read_text(encoding='utf-8-sig')
    except UnicodeError: raise ConversionError('Text/HTML must use UTF-8 encoding.', 'invalid_encoding')
    if '\x00' in text or any(ord(c)<32 and c not in '\n\r\t' for c in text): raise ConversionError('Binary content is not accepted as text.', 'invalid_file')
    if len(text)>settings.max_text_chars: raise ConversionError('Text exceeds the configured document limit.', 'text_limit')
    if not text.strip(): raise ConversionError('The document is empty.', 'empty_document')
    return text


def docx_html(path):
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    doc=Document(path); parts=[]
    for child in doc.element.body.iterchildren():
        if child.tag.endswith('}p'):
            p=Paragraph(child,doc); runs=[]
            for run in p.runs:
                value=escape(run.text).replace('\n','<br>')
                if run.bold: value='<strong>'+value+'</strong>'
                if run.italic: value='<em>'+value+'</em>'
                if run.underline: value='<u>'+value+'</u>'
                runs.append(value)
            # python-docx includes hyperlink text in p.text but not p.runs in some versions.
            value=''.join(runs) if ''.join(r.text for r in p.runs)==p.text else escape(p.text)
            tag='h2' if p.style.name.startswith('Heading') else 'p'
            if p.text.strip(): parts.append(f'<{tag}>'+value+f'</{tag}>')
        elif child.tag.endswith('}tbl'):
            for row in Table(child,doc).rows: parts.append('<p>'+escape(' | '.join(c.text for c in row.cells))+'</p>')
    for section in doc.sections:
        for p in list(section.header.paragraphs)+list(section.footer.paragraphs):
            if p.text.strip(): parts.append('<p>'+escape(p.text)+'</p>')
    html=clean_html(''.join(parts))
    if not plain_text(html): raise ConversionError('The DOCX has no readable text.', 'empty_document')
    return html


def open_pdf(path):
    try:
        pdf=pymupdf.open(path)
        if pdf.needs_pass: pdf.close(); raise ConversionError('Password-protected PDFs require an unlocked copy.', 'password_protected')
        if not 1 <= len(pdf) <= settings.max_pages: pdf.close(); raise ConversionError(f'PDFs must contain 1–{settings.max_pages} pages.', 'page_limit')
        if any(page.rect.width*page.rect.height>4_000_000 for page in pdf): pdf.close(); raise ConversionError('PDF page dimensions exceed the processing limit.', 'page_limit')
        return pdf
    except ConversionError: raise
    except Exception: raise ConversionError('The PDF is corrupted or unreadable.', 'corrupted_document')


def _docx_compact_styles(doc, body_size, leading, font_name):
    """Replace python-docx defaults (10 pt after every paragraph, 1.15 line height,
    24 pt before headings) that made converted pages spill onto extra pages."""
    normal=doc.styles['Normal'];normal.font.size=Pt(body_size);normal.font.name=font_name
    fmt=normal.paragraph_format;fmt.space_before=Pt(0);fmt.space_after=Pt(round(body_size*.35,1));fmt.line_spacing=Pt(leading);fmt.widow_control=True
    for name,size,before in (('Heading 1',body_size*1.55,body_size*.9),('Heading 2',body_size*1.15,body_size*.9)):
        style=doc.styles[name];style.font.size=Pt(round(size,1));style.font.name=font_name
        style.paragraph_format.space_before=Pt(round(before,1));style.paragraph_format.space_after=Pt(round(body_size*.3,1));style.paragraph_format.line_spacing=Pt(round(size*1.2,1));style.paragraph_format.keep_with_next=True
    bullet=doc.styles['List Bullet'];bullet.paragraph_format.space_after=Pt(round(body_size*.2,1));bullet.paragraph_format.line_spacing=Pt(leading)


def _source_metrics(pdf):
    """Page size, margins and dominant body type from the source PDF text layer."""
    sizes=Counter();fonts=Counter();lefts=[];rights=[];tops=[];bottoms=[]
    for page in pdf:
        blocks=[b for b in page.get_text('dict')['blocks'] if b.get('type')==0]
        if not blocks: continue
        lefts.append(min(b['bbox'][0] for b in blocks));rights.append(page.rect.width-max(b['bbox'][2] for b in blocks))
        tops.append(min(b['bbox'][1] for b in blocks));bottoms.append(page.rect.height-max(b['bbox'][3] for b in blocks))
        for b in blocks:
            for line in b['lines']:
                for span in line['spans']:
                    if span['text'].strip(): sizes[round(span['size']*2)/2]+=len(span['text']);fonts[span.get('font','')]+=len(span['text'])
    body=sizes.most_common(1)[0][0] if sizes else 10.5
    dominant=(fonts.most_common(1)[0][0] if fonts else '').lower()
    family='Noto Serif' if any(v in dominant for v in ('times','serif','georgia','garamond','cambria')) and 'sans' not in dominant else 'Noto Sans Mono' if any(v in dominant for v in ('mono','courier','consol')) else 'Noto Sans'
    clamp=lambda values,default:max(18,min(72,sorted(values)[len(values)//2])) if values else default
    return {'body':max(8,min(13,body)),'font':family,'left':clamp(lefts,54),'right':clamp(rights,54),'top':clamp(tops,48),'bottom':clamp(bottoms,48)}


def _add_runs(paragraph, html):
    """Insert the safe inline HTML produced by pdf_layout (strong/em only) as runs."""
    for token in re.split(r'(<strong>.*?</strong>|<em>.*?</em>)', html):
        if not token: continue
        bold=token.startswith('<strong>');italic=token.startswith('<em>') or '<em>' in token
        text=unescape(re.sub(r'<[^>]+>','',token))
        if text:
            run=paragraph.add_run(text);run.bold=bold or None;run.italic=italic or None


def pdf_to_docx(path, output):
    """Editable Word reconstruction in reading order.

    Uses the column-aware layout engine (sidebars are read one column at a time,
    wrapped lines become real paragraphs, bullets become Word list items, page
    numbers/running headers are dropped). Page size, margins, body type size and
    font family are taken from the source so content does not spill onto extra
    pages; source page boundaries are kept as page breaks only when needed.
    """
    warnings=['PDF-to-Word reconstructs editable content. Columns become one reading stream; diagrams, images, signatures and exact positions are not reproduced, so this is not an exact replica.']
    from .resume_import import heading as section_heading
    from .pdf_layout import analyze_document
    doc=Document();found=False
    with open_pdf(path) as pdf:
        metrics=_source_metrics(pdf);first=pdf[0].rect
        section=doc.sections[0];section.page_width,section.page_height=Pt(first.width),Pt(first.height)
        section.left_margin,section.right_margin=Pt(metrics['left']),Pt(metrics['right']);section.top_margin,section.bottom_margin=Pt(metrics['top']),Pt(metrics['bottom'])
        body=metrics['body'];leading=round(body*1.28,1)
        _docx_compact_styles(doc,body,leading,metrics['font'])
        try: layouts={layout.number:layout for layout in analyze_document(pdf,lambda value:section_heading(value) is not None)}
        except Exception: layouts={};warnings.append('Layout analysis was limited; text follows spatial order.')
        removed=[]
        for number,page in enumerate(pdf):
            text=page.get_text('text',sort=True)
            if len(text.strip())<30 and page.get_images():
                if number and found: doc.add_page_break()
                try: text=ocr_page(page)
                except ExtractionError as exc: raise ConversionError(str(exc),'ocr_unavailable')
                warnings.append(f'OCR used on page {number+1}; verify recognition accuracy.')
                for p in text.split('\n\n'):
                    if p.strip(): doc.add_paragraph(p.strip()); found=True
                continue
            tables=[]
            try:
                for table in page.find_tables().tables:
                    rows=table.extract()
                    if rows and len(rows)>=2 and len(rows[0])>=2: tables.append((pymupdf.Rect(table.bbox),rows))
            except Exception: warnings.append(f'Table reconstruction was limited on page {number+1}.')
            layout=layouts.get(number+1)
            if layout is None:
                for line in text.splitlines():
                    if line.strip(): doc.add_paragraph(line.strip()); found=True
                continue
            removed.extend(layout.removed)
            placed=set()
            for para in layout.paragraphs:
                inside=next((index for index,(rect,_) in enumerate(tables) if rect.y0-1<=para.y0<=rect.y1+1 and rect.x0-1<=para.x0<=rect.x1+1),None)
                if inside is not None:
                    if inside not in placed:
                        placed.add(inside);rows=tables[inside][1];built=doc.add_table(rows=0,cols=len(rows[0]));built.style='Table Grid'
                        for row in rows:
                            cells=built.add_row().cells
                            for index,value in enumerate(row): cells[index].text=value or ''
                        found=True
                    continue
                if para.bullet:
                    p=doc.add_paragraph(style='List Bullet')
                elif section_heading(para.text):
                    p=doc.add_paragraph(style='Heading 2')
                elif para.size>=body*1.8 and len(para.text)<80 and not found:
                    p=doc.add_paragraph(style='Heading 1')
                else:
                    p=doc.add_paragraph()
                    if para.size<body-1: 
                        p.paragraph_format.space_after=Pt(round(body*.2,1))
                _add_runs(p,para.html)
                if not p.style.name.startswith('Heading') and abs(para.size-body)>1.2:
                    for run in p.runs: run.font.size=Pt(round(para.size*2)/2)
                p.paragraph_format.keep_together=True
                found=True
            # A source page break is kept only if the next page starts a new section;
            # otherwise the flow continues so short overflow does not leave half-empty pages.
            following=layouts.get(number+2)
            if following and following.paragraphs and section_heading(following.paragraphs[0].text) and len(layout.paragraphs)>12:
                doc.add_page_break()
        if removed: warnings.append(f'{len(removed)} page-number or repeated header line(s) were left out.')
        if any(layout.columns==2 for layout in layouts.values()): warnings.append('A column or sidebar layout was converted into one reading stream (left column first).')
    if not found: raise ConversionError('No editable text could be extracted from the PDF.', 'empty_document')
    doc.save(output)
    return list(dict.fromkeys(warnings))


def image(path):
    try:
        source=Image.open(path)
        if source.width*source.height>20_000_000: raise ConversionError('Image exceeds 20 megapixels.', 'image_limit')
        source.load()
        return ImageOps.exif_transpose(source)
    except ConversionError: raise
    except Exception: raise ConversionError('Image is corrupted or exceeds safe decoding limits.', 'corrupted_image')


def rgb(source):
    if source.mode in ('RGBA','LA') or 'transparency' in source.info:
        alpha=source.convert('RGBA'); background=Image.new('RGB',source.size,'white');background.paste(alpha,mask=alpha.getchannel('A'));return background
    return source.convert('RGB')


def images_pdf(paths, output, options):
    width,height=(595.28,841.89) if options.page_size=='A4' else (612,792)
    if options.orientation=='landscape': width,height=height,width
    margin=options.margin_mm*72/25.4
    doc=canvas.Canvas(str(output),pagesize=(width,height))
    for path in paths:
        img=rgb(image(path));scale=min((width-2*margin)/img.width,(height-2*margin)/img.height)
        w,h=img.width*scale,img.height*scale
        doc.drawImage(ImageReader(img),(width-w)/2,(height-h)/2,width=w,height=h,preserveAspectRatio=True)
        doc.showPage();img.close()
    doc.save()


def office_pdf(path, work):
    if not office_available(): raise ConversionError('Legacy Office conversion requires enabled LibreOffice and bubblewrap. Upload DOCX or export a PDF first.', 'utility_unavailable')
    # No network namespace; bind only system binaries/fonts and this job workspace.
    command=['bwrap','--unshare-all','--die-with-parent','--new-session','--clearenv','--ro-bind','/usr','/usr','--proc','/proc','--dev','/dev','--tmpfs','/tmp']
    for directory in ['/lib','/lib64','/bin']:
        if Path(directory).exists(): command+=['--ro-bind',directory,directory]
    if Path('/etc/fonts').exists(): command+=['--ro-bind','/etc/fonts','/etc/fonts']
    command+=['--bind',str(work.resolve()),'/work','--chdir','/work','--setenv','HOME','/tmp','--setenv','PATH','/usr/bin:/bin','--setenv','LANG','C.UTF-8','--',shutil.which('libreoffice'),'-env:UserInstallation=file:///tmp/lo-profile','--headless','--norestore','--nodefault','--nolockcheck','--convert-to','pdf','--outdir','/work', '/work/'+path.name]
    completed=subprocess.run(command,capture_output=True,timeout=60,check=False)
    output=work/(path.stem+'.pdf')
    if completed.returncode or not output.exists(): raise ConversionError('The sandboxed Office engine failed. Check namespace support or use PDF/DOCX.', 'office_failed')
    return output


def validate_output(path, ext):
    if not path.is_file() or path.stat().st_size==0: raise ConversionError('No output file was produced.', 'empty_output')
    if path.stat().st_size>100*1024*1024: raise ConversionError('Output exceeds the configured 100 MB limit.', 'output_limit')
    try:
        if ext=='pdf':
            with pymupdf.open(path) as pdf:
                if not len(pdf) or pdf.needs_pass: raise ValueError()
        elif ext=='docx': Document(path)
        elif ext in IMAGE_TYPES:
            with Image.open(path) as img: img.verify()
        elif ext in ('txt','html'):
            if not path.read_text(encoding='utf-8').strip(): raise ValueError()
        elif ext=='zip':
            with zipfile.ZipFile(path) as z:
                if not z.namelist() or z.testzip(): raise ValueError()
    except Exception: raise ConversionError('The generated output could not be opened and validated.', 'invalid_output')


def convert(paths:list[Path], source:str, target:str, work:Path, options:ConversionSettings):
    if source not in matrix() or target not in matrix()[source]: raise ConversionError('This source/target pair is not supported by this backend.', 'unsupported_conversion')
    if len(paths)>1 and not (source in IMAGE_TYPES and target=='pdf'): raise ConversionError('Multiple inputs are only supported for images to one PDF.', 'unsupported_batch')
    outputs=[];warnings=[];path=paths[0]
    output=work/('converted.'+target)
    if source in IMAGE_TYPES:
        if target=='pdf': images_pdf(paths,output,options)
        else:
            img=image(path)
            if target in ('jpg','webp'): img=rgb(img)
            img.save(output,format=IMAGE_TYPES[target],quality=90);img.close()
    elif source=='pdf':
        if target in ('jpg','png'):
            with open_pdf(path) as pdf:
                for index in page_indices(options.pages,len(pdf)):
                    page=pdf[index];scale=options.dpi/72
                    if page.rect.width*page.rect.height*scale*scale>30_000_000: raise ConversionError('Rendered page exceeds the pixel limit. Choose a lower resolution.', 'image_limit')
                    output=work/f'page-{index+1:03d}.{target}';page.get_pixmap(dpi=options.dpi,alpha=False).save(output);outputs.append(output)
            archive=work/'all-pages.zip'
            with zipfile.ZipFile(archive,'w',compression=zipfile.ZIP_DEFLATED) as z:
                for item in outputs: z.write(item,item.name)
            outputs.append(archive)
        elif target=='docx': warnings=pdf_to_docx(path,output)
        else:
            try: result=extract(path,'pdf')
            except ExtractionError as exc: raise ConversionError(str(exc),'extraction_failed')
            output.write_text(result['text'],encoding='utf-8');warnings=[w['message'] for w in result['extraction']['warnings']]
    elif source=='docx':
        html=docx_html(path) if target!='pdf' else ''
        if target=='pdf': output=office_pdf(path,work)
        elif target=='txt': output.write_text(plain_text(html),encoding='utf-8')
        elif target=='html': output.write_text('<!doctype html><html lang="en"><meta charset="utf-8"><title>Converted document</title><body>'+html+'</body></html>',encoding='utf-8')
    elif source in ('txt','html'):
        text=safe_text(path)
        if source=='html':
            text=re.sub(r'<(script|style)[^>]*>.*?</\1\s*>','',text,flags=re.I|re.S)
            html=clean_html(text)
            text=plain_text(html)
            if not text.strip(): raise ConversionError('The HTML contains no readable text.', 'empty_document')
            warnings=['HTML is sanitized and reconstructed without CSS, scripts, remote images or network requests.']
        else: html=text_to_html(text)
        if target=='txt': output.write_text(text,encoding='utf-8')
        else:
            content=DraftContent(sections=[Section(id='document',title='Contact information',html=html)])
            output.write_bytes(pdf_bytes(content) if target=='pdf' else docx_bytes(content))
    else:
        output=office_pdf(path,work)
    if not outputs: outputs=[output]
    result=[]
    for item in outputs:
        ext=item.suffix.lstrip('.');validate_output(item,ext)
        result.append({'filename':item.name,'format':ext,'size':item.stat().st_size,'mime':MIMES[ext]})
    if sum(o['size'] for o in result)>150*1024*1024: raise ConversionError('Combined output exceeds 150 MB. Convert fewer pages.', 'output_limit')
    return {'outputs':result,'warnings':warnings}
