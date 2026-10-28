"""Native Word approximation of supported visual layouts, never a renamed PDF.

Word supports margins, headings, emphasis, lists, header arrangements, column
container tables and page breaks. Browser pagination and full-bleed treatments
are not identical; this limitation is disclosed in the studio.
"""
import io
import re
from docx import Document
from docx.shared import Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .editor_content import DraftContent
from .document_renderer import effective_definition
from .exports import blocks,hyperlink


def shade(p,color):
    node=OxmlElement('w:shd');node.set(qn('w:fill'),color.lstrip('#'));p._p.get_or_add_pPr().append(node)


def border(p,color,side='bottom'):
    borders=OxmlElement('w:pBdr');line=OxmlElement('w:'+side);line.set(qn('w:val'),'single');line.set(qn('w:sz'),'6');line.set(qn('w:color'),color.lstrip('#'));borders.append(line);p._p.get_or_add_pPr().append(borders)


def zero_cell_padding(table):
    """Column/label tables are layout containers; Word's default cell padding and
    autofit would otherwise widen content and push it onto extra pages."""
    tbl_pr=table._tbl.tblPr
    margins=OxmlElement('w:tblCellMar')
    for side,value in (('top',0),('left',0),('bottom',0),('right',170)):
        node=OxmlElement('w:'+side);node.set(qn('w:w'),str(value));node.set(qn('w:type'),'dxa');margins.append(node)
    tbl_pr.append(margins)
    layout=OxmlElement('w:tblLayout');layout.set(qn('w:type'),'fixed');tbl_pr.append(layout)


def spacing(paragraph_format,before,after):
    paragraph_format.space_before=Pt(before);paragraph_format.space_after=Pt(after)


def studio_docx_bytes(content:DraftContent,fit:dict|None=None):
    """Word file whose vertical rhythm matches the studio canvas.

    python-docx's default template adds 24 pt before every Heading 1, 10 pt after
    every paragraph and 1.15 line height. Those defaults pushed one-page resumes onto
    a second page, so every style used here sets explicit spacing.

    ``fit`` is the bounded tightening step the canvas applied to keep a short overflow
    on fewer pages (see pagination.js FIT_STEPS); the same scales are applied here so
    PDF and Word agree on page count.
    """
    cfg=dict(effective_definition(content));doc=Document();page=doc.sections[0]
    if fit:
        cfg['spacing']=cfg['spacing']*float(fit.get('space',1));cfg['leading']=cfg['leading']*float(fit.get('leading',1))
        cfg['body_size']=max(9.0,round(cfg['body_size']*float(fit.get('body',1))*2)/2);cfg['margin']=max(30.0,cfg['margin']*float(fit.get('margin',1)))
        cfg['leading']=max(cfg['leading'],cfg['body_size']*1.12)
    w,h=(595.28,841.89) if content.style.page_size=='A4' else (612,792)
    page.page_width,page.page_height=Pt(w),Pt(h)
    page.left_margin=page.right_margin=page.top_margin=page.bottom_margin=Pt(cfg['margin'])
    page.header_distance=page.footer_distance=Pt(min(24,cfg['margin']/2))
    leading=Pt(max(cfg['leading'],cfg['body_size']*1.1));gap=cfg['spacing']*.6
    normal=doc.styles['Normal'];normal.font.name={'sans':'Noto Sans','serif':'Noto Serif','mono':'Noto Sans Mono'}[cfg['font']];normal.font.size=Pt(cfg['body_size']);normal.paragraph_format.line_spacing=leading;spacing(normal.paragraph_format,0,gap)
    normal.paragraph_format.widow_control=True
    heading=doc.styles['Heading 1'];heading.font.name=normal.font.name;heading.font.size=Pt(cfg['heading_size']);heading.font.color.rgb=RGBColor.from_string(cfg['accent'].lstrip('#'));heading.paragraph_format.keep_with_next=True
    heading.paragraph_format.line_spacing=Pt(cfg['heading_size']*1.25);spacing(heading.paragraph_format,cfg['spacing']*1.6,cfg['spacing']*.7)
    bullet_style=doc.styles['List Bullet'];bullet_style.paragraph_format.line_spacing=leading;spacing(bullet_style.paragraph_format,0,gap*.6)
    numbering=doc.part.numbering_part.element
    abstract_id=max([int(n.get(qn('w:abstractNumId'))) for n in numbering.findall(qn('w:abstractNum'))]+[0])+1
    num_id=max([int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))]+[0])+1
    abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),str(abstract_id));level=OxmlElement('w:lvl');level.set(qn('w:ilvl'),'0')
    for tag,value in [('start','1'),('numFmt','bullet'),('lvlText',{'disc':'•','square':'▪','dash':'–'}[cfg['bullet']])]:
        element=OxmlElement('w:'+tag);element.set(qn('w:val'),value);level.append(element)
    abstract.append(level);numbering.append(abstract);num=OxmlElement('w:num');num.set(qn('w:numId'),str(num_id));reference=OxmlElement('w:abstractNumId');reference.set(qn('w:val'),str(abstract_id));num.append(reference);numbering.append(num)
    def paragraph(container,block,contact=False,index=0):
        p=container.add_paragraph(style='List Bullet' if block['bullet'] else None)
        if block['bullet']:
            properties=p._p.get_or_add_pPr().get_or_add_numPr();properties.get_or_add_ilvl().val=0;properties.get_or_add_numId().val=num_id
            p.paragraph_format.left_indent=Pt(14);p.paragraph_format.first_line_indent=Pt(-10)
        if contact:
            spacing(p.paragraph_format,0,cfg['spacing']*.35 if index else cfg['spacing']*.5)
            if index==0:p.paragraph_format.line_spacing=Pt(cfg['name_size']*1.15)
        p.paragraph_format.keep_together=True
        if contact and cfg['header']=='center':p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if content.style.alignment=='justify' and not contact:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        text=''.join(value['text'] for value in block['runs'])
        if not contact and content.style.date_alignment=='right' and len(text)<110 and re.search(r'\b(?:19|20)\d{2}\b',text) and not any(value['bold'] for value in block['runs']):p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        for value in block['runs']:
            if value['href']:hyperlink(p,value)
            else:
                run=p.add_run(value['text']);run.bold=value['bold'];run.italic=value['italic'];run.underline=value['underline']
                if contact and index==0:run.font.size=Pt(cfg['name_size']);run.bold=True
                if contact and cfg['header']=='band':run.font.color.rgb=RGBColor(255,255,255)
        if contact and cfg['header']=='band':shade(p,cfg['accent'])
        if not contact and cfg['entries']=='timeline':border(p,'#cbd2d9','left')
        if not contact and cfg['entries']=='card':shade(p,'#f5f6f8')
        if not contact and cfg['entries']=='indented' and block['bullet']:p.paragraph_format.left_indent=Pt(12)
        return p
    def render_section(section,container,number):
        target=container
        if section.page_break_before:
            p=container.add_paragraph();p.paragraph_format.page_break_before=True
        if cfg['layout']=='labelled':
            table=container.add_table(rows=1,cols=2);table.autofit=False;zero_cell_padding(table)
            widths=(Pt((w-2*cfg['margin'])*.25),Pt((w-2*cfg['margin'])*.75))
            for column,width in zip(table.columns,widths):column.width=width
            for cell,width in zip(table.rows[0].cells,widths):cell.width=width
            label=table.cell(0,0).paragraphs[0];label.text=section.title.upper();label.style='Heading 1';spacing(label.paragraph_format,cfg['spacing'],0)
            target=table.cell(0,1);target._tc.remove(target.paragraphs[0]._p)
            for run in label.runs:run.font.size=Pt(max(8,cfg['heading_size']*.8))
        else:
            p=container.add_paragraph(style='Heading 1');p.text=(str(number).zfill(2)+'  ' if cfg['heading']=='numbered' else '')+(section.title.upper() if cfg['heading']=='caps' else section.title)
            if cfg['dividers'] and cfg['heading'] in ('rule','underline'):border(p,cfg['accent'])
            if cfg['heading']=='bar':shade(p,'#eef0f7');border(p,cfg['accent'],'left')
            if cfg['heading']=='editorial':p.alignment=WD_ALIGN_PARAGRAPH.CENTER;[setattr(run,'italic',True) for run in p.runs]
        rendered=[paragraph(target,block) for block in blocks(section.html)]
        if cfg['layout']=='labelled' and rendered:spacing(rendered[0].paragraph_format,cfg['spacing'],gap)
        if cfg['layout']=='labelled' and not rendered:target.add_paragraph()
    contact=next((s for s in content.sections if s.title.lower()=='contact information'),None)
    others=[s for s in content.sections if s is not contact]
    def kind(s):return next((k for k in ['skills','education','certifications','projects','summary','languages','experience'] if k in s.title.lower()),'custom')
    def contact_into(container):
        if not contact:return
        bs=blocks(contact.html)
        if cfg['header']=='split' and len(bs)>1:
            table=container.add_table(rows=1,cols=2);table.autofit=False;zero_cell_padding(table)
            half=Pt((w-2*cfg['margin'])/2)
            for cell in table.rows[0].cells:cell.width=half
            for index,block in enumerate(bs):paragraph(table.cell(0,0 if index<2 else 1),block,True,index)
            for cell in table.rows[0].cells:
                if len(cell.paragraphs)>1 and not cell.paragraphs[0].text:cell._tc.remove(cell.paragraphs[0]._p)
            for p in table.cell(0,1).paragraphs:p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        else:
            for index,block in enumerate(bs):paragraph(container,block,True,index)
    if cfg['header']!='rail':contact_into(doc)
    lead=[s for s in others if kind(s) in cfg['lead'] and not s.page_break_before]
    rest=[s for s in others if s not in lead]
    for index,section in enumerate(lead):render_section(section,doc,index+1)
    if cfg['columns']==2:
        table=doc.add_table(rows=1,cols=2);table.autofit=False;zero_cell_padding(table)
        rail_index=0 if cfg['layout']=='sidebar-left' else 1
        main_index=1-rail_index;ratio=cfg['column_ratio']
        widths={main_index:Pt((w-2*cfg['margin'])*ratio),rail_index:Pt((w-2*cfg['margin'])*(1-ratio))}
        for index,column in enumerate(table.columns):column.width=widths[index]
        for index,cell in enumerate(table.rows[0].cells):cell.width=widths[index]
        rail_cell,main_cell=table.cell(0,rail_index),table.cell(0,main_index)
        if cfg['header']=='rail':contact_into(rail_cell)
        for index,section in enumerate(rest):render_section(section,rail_cell if section.placement=='rail' or section.placement!='main' and kind(section) in cfg['rail'] else main_cell,index+1)
        for cell in (rail_cell,main_cell):
            # Drop the empty paragraph Word creates in every new cell.
            if len(cell.paragraphs)>1 and not cell.paragraphs[0].text and not cell.paragraphs[0]._p.xpath('./w:pPr/w:numPr'):cell._tc.remove(cell.paragraphs[0]._p)
    else:
        for index,section in enumerate(rest):render_section(section,doc,index+1)
    first=next((p for p in doc.paragraphs if p.style.name=='Heading 1'),None)
    if first is not None and not contact:spacing(first.paragraph_format,0,cfg['spacing']*.7)
    buffer=io.BytesIO();doc.save(buffer);data=buffer.getvalue();Document(io.BytesIO(data))
    return data


async def measured_docx_bytes(content:DraftContent):
    """DOCX using the page-fit the real PDF layout needed. Falls back to the template
    spacing when Chromium is unavailable (the DOCX is still valid)."""
    import asyncio
    fit=None
    if content.style.auto_fit:
        try:
            from .document_renderer import render_pdf
            _,info=await asyncio.wait_for(render_pdf(content),timeout=65);fit=info.get('fit')
        except Exception:
            fit=None
    return await asyncio.to_thread(studio_docx_bytes,content,fit)
