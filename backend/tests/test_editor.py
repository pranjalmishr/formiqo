import io
import json
import asyncio
import pytest
import pymupdf
from docx import Document
from fastapi.testclient import TestClient
from app.main import app, expire_records
from app.db import Session, Resume, EditorDraft, EditorVersion, now
from app.editor_content import DraftContent, RewriteRequest, clean_html, plain_text
from app.editor_api import request_rewrite
from app.config import settings
from datetime import timedelta
from sqlalchemy import select


def open_draft(client, resume_text):
    uploaded = client.post('/api/resumes/upload', files={'file': ('resume.txt', resume_text.encode(), 'text/plain')})
    assert uploaded.status_code == 202
    rid = uploaded.json()['id']
    result = client.get('/api/editor/'+rid)
    assert result.status_code == 200, result.text
    return result.json()


def save(client, draft, **extra):
    return client.patch('/api/editor/'+draft['resume_id'], json={'revision': draft['revision'], 'content': draft['content'], 'applied': draft['applied'], 'rejected': draft['rejected'], **extra})


def test_seed_edit_save_versions_and_original(client, resume_text):
    d = open_draft(client, resume_text)
    assert d['revision'] == 1 and len(d['content']['sections']) >= 4
    d['content']['sections'][1]['html'] = '<p><strong>Verified</strong> professional summary.</p>'
    response = save(client, d)
    assert response.status_code == 200, response.text
    updated = response.json()
    assert updated['revision'] == 2
    assert '<strong>Verified</strong>' in updated['content']['sections'][1]['html']
    r = client.get('/api/resumes/'+d['resume_id']).json()
    assert r['original_text'] == d['original_text'] and 'Verified' in r['text']
    versions = client.get('/api/editor/'+d['resume_id']+'/versions').json()['items']
    assert len(versions) == 2 and versions[0]['revision'] == 2
    # Restoring an older content snapshot creates a new revision; no history is destroyed.
    updated['content'] = versions[-1]['content']
    restored = save(client, updated).json()
    assert restored['revision'] == 3


def test_revision_conflict_and_external_preview_conflict(client, resume_text):
    d = open_draft(client, resume_text)
    d['content']['sections'][0]['html'] += '<p>Verified link</p>'
    assert save(client, d).status_code == 200
    assert save(client, d).status_code == 409
    client.patch('/api/resumes/'+d['resume_id'], json={'text': resume_text+'\nNew verified project'})
    current = client.get('/api/editor/'+d['resume_id']).json()
    assert current['out_of_sync']
    assert save(client, current).status_code == 409
    imported = client.post('/api/editor/'+d['resume_id']+'/import-preview', json={'revision': current['revision'], 'content': current['content']})
    assert imported.status_code == 200 and not imported.json()['out_of_sync']
    assert 'New verified project' in str(imported.json()['content'])


def test_sanitization_and_bullet_text():
    clean = clean_html('<script>alert(1)</script><p onclick="run()">Hello <strong>world</strong><img src=x><a href="javascript:alert(1)">link</a></p>')
    assert 'script' not in clean and 'onclick' not in clean and '<img' not in clean and 'javascript' not in clean
    assert '<strong>world</strong>' in clean
    assert '• Built verified services' in plain_text('<ul><li><p>Built verified services</p></li></ul>')
    with pytest.raises(ValueError):
        DraftContent.model_validate({'sections':[{'id':'a','title':'One','html':'<p>A</p>'},{'id':'a','title':'Two','html':'<p>B</p>'}]})


def test_exports_are_real_readable_and_ordered(client, resume_text):
    d = open_draft(client, resume_text)
    d['content']['sections'] = [
        {'id':'contact','title':'Contact information','html':'<p>Jordan Rivera</p><p><a href="mailto:jordan@example.com">jordan@example.com</a></p>'},
        {'id':'projects','title':'Projects','html':'<ul><li><p><strong>Built</strong> verified Python services.</p></li></ul>'},
        {'id':'education','title':'Education','html':'<p>Actual coursework</p>'},
    ]
    assert save(client,d).status_code == 200
    pdf = client.get('/api/editor/'+d['resume_id']+'/export?format=pdf')
    assert pdf.status_code == 200, pdf.text[:100]
    assert pdf.content.startswith(b'%PDF-')
    with pymupdf.open(stream=pdf.content,filetype='pdf') as doc:
        text='\n'.join(page.get_text() for page in doc)
        assert 'Jordan Rivera' in text and text.index('Projects') < text.index('Education')
        assert any(page.get_links() for page in doc)
    docx = client.get('/api/editor/'+d['resume_id']+'/export?format=docx')
    assert docx.status_code == 200
    doc=Document(io.BytesIO(docx.content))
    paragraphs=[p.text for p in doc.paragraphs]
    assert paragraphs.index('Projects') < paragraphs.index('Education')
    assert any(p.style.name=='List Bullet' for p in doc.paragraphs)
    assert any('mailto:' in rel.target_ref for rel in doc.part.rels.values() if rel.is_external)
    assert client.get('/api/editor/'+d['resume_id']+'/export?format=exe').status_code == 422


def test_editor_ownership_and_deletion(client, resume_text):
    d=open_draft(client,resume_text)
    with TestClient(app) as other:
        other.get('/api/config')
        for suffix in ['', '/versions', '/export?format=pdf']:
            assert other.get('/api/editor/'+d['resume_id']+suffix).status_code == 404
        assert save(other,d).status_code == 404
    assert client.delete('/api/resumes/'+d['resume_id']).status_code == 204
    with Session() as db:
        assert db.get(EditorDraft,d['resume_id']) is None
        assert not db.scalars(select(EditorVersion).where(EditorVersion.resume_id==d['resume_id'])).all()


def test_editor_expiration_cleanup(client,resume_text):
    d=open_draft(client,resume_text)
    with Session() as db:
        r=db.get(Resume,d['resume_id']);r.expires_at=now()-timedelta(minutes=1);db.commit()
    expire_records()
    assert client.get('/api/editor/'+d['resume_id']).status_code==404
    with Session() as db:
        assert db.get(EditorDraft,d['resume_id']) is None


def rewrite_payload(d):
    s=d['content']['sections'][1]
    return {'resume_id':d['resume_id'],'section_id':s['id'],'revision':d['revision'],'original':plain_text(s['html']),'action':'Improve wording','ai_consent':True}


def test_ai_consent_configuration_and_source_validation(client,resume_text,monkeypatch):
    d=open_draft(client,resume_text);payload=rewrite_payload(d)
    assert client.post('/api/ai/improve',json={**payload,'ai_consent':False}).status_code==403
    monkeypatch.setattr(settings,'llm_api_key','')
    assert client.post('/api/ai/improve',json=payload).status_code==503
    assert client.post('/api/ai/improve',json={**payload,'original':'Invented source'}).status_code==422
    assert client.post('/api/ai/improve',json={**payload,'revision':999}).status_code==409


def test_ai_schema_no_numbers_and_no_automatic_mutation(client,resume_text,monkeypatch):
    d=open_draft(client,resume_text);payload=rewrite_payload(d)
    monkeypatch.setattr(settings,'llm_api_key','testing-only')
    output={'original':payload['original'],'suggested':'Engineering tools with [actual outcome].','notes':['Replace the placeholder with your verified outcome.']}
    class FakeResponse:
        def raise_for_status(self): pass
        def json(self):return {'choices':[{'message':{'content':json.dumps(output)}}]}
    class FakeClient:
        def __init__(self,**kwargs):pass
        async def __aenter__(self):return self
        async def __aexit__(self,*args):pass
        async def post(self,*args,**kwargs):return FakeResponse()
    monkeypatch.setattr('app.editor_api.httpx.AsyncClient',FakeClient)
    result=client.post('/api/ai/improve',json=payload)
    assert result.status_code==200 and result.json()['requires_confirmation']
    assert client.get('/api/editor/'+d['resume_id']).json()['content']==d['content']
    output['suggested']='Generated 987654 percent growth.'
    assert client.post('/api/ai/improve',json=payload).status_code==502
    output['suggested']='Safe text';output['score']=99
    assert client.post('/api/ai/improve',json=payload).status_code==502


def test_applied_rejected_tracking_without_score_bonus(client,resume_text):
    d=open_draft(client,resume_text)
    before=client.get('/api/resumes/'+d['resume_id']).json()['text']
    updated=save(client,d,applied=['confirmed-suggestion'],rejected=['discarded-suggestion']).json()
    assert updated['applied']==['confirmed-suggestion'] and updated['rejected']==['discarded-suggestion']
    from app.scoring import analyze
    after=client.get('/api/resumes/'+d['resume_id']).json()['text']
    assert analyze(before,{})['overall_score']==analyze(after,{})['overall_score']
