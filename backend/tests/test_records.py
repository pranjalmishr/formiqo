"""Score records: listing, filters, ownership and real CSV/XLSX/JSON exports."""
import csv
import io
import json
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.main import app

JD = 'Requirements: Python, PostgreSQL and AWS experience. Responsibilities: build reliable services. Nice to have: Kubernetes.'


def ready_resume(client, text, name='resume.txt'):
    response = client.post('/api/resumes/upload', files={'file': (name, text.encode(), 'text/plain')})
    assert response.status_code == 202, response.text
    return client.get('/api/resumes/' + response.json()['id']).json()


def test_records_list_filter_and_exports(client, resume_text):
    resume = ready_resume(client, resume_text, 'jordan.txt')
    general = client.post('/api/ats/analyze', json={'resume_id': resume['id'], 'text': resume['text']}).json()
    matched = client.post('/api/ai/match-job-description', json={'resume_id': resume['id'], 'job_description': JD, 'job_title': 'Backend Engineer', 'company': 'Acme'}).json()
    listing = client.get('/api/records').json()
    assert listing['total'] == 2 and listing['summary']['records'] == 2 and listing['summary']['job_matches'] == 1
    rows = {row['analysis_id']: row for row in listing['items']}
    assert rows[general['id']]['ats_score'] == general['overall_score'] and rows[general['id']]['job_match_score'] is None
    job = rows[matched['id']]
    assert job['mode'] == 'Job match' and job['job_match_score'] == matched['match_score']
    assert job['job_title'] == 'Backend Engineer' and job['company'] == 'Acme'
    assert job['required_terms_total'] >= 1 and job['required_terms_found'] <= job['required_terms_total']
    assert job['weighted_term_coverage'] is not None and 0 <= job['weighted_term_coverage'] <= 100
    assert client.get('/api/records?mode=job').json()['total'] == 1
    assert client.get('/api/records?q=acme').json()['total'] == 1
    assert client.get('/api/records?q=nothing-like-this').json()['total'] == 0
    # CSV opens as a spreadsheet with one row per analysis.
    response = client.get('/api/records/export?format=csv')
    assert response.status_code == 200 and 'attachment' in response.headers['content-disposition']
    table = list(csv.DictReader(io.StringIO(response.content.decode('utf-8-sig'))))
    assert len(table) == 2 and {row['analysis_id'] for row in table} == {general['id'], matched['id']}
    assert next(row for row in table if row['analysis_id'] == matched['id'])['job_match_score'] == str(matched['match_score'])
    # XLSX is a real workbook with records + summary sheets.
    response = client.get('/api/records/export?format=xlsx')
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames[:2] == ['Score records', 'Summary'] and book['Score records'].max_row == 3
    # JSON keeps full evidence for each analysis.
    full = json.loads(client.get('/api/records/export?format=json').content)
    assert len(full['records']) == 2 and any(item['keyword_terms'] for item in full['records'])


def test_records_are_private_and_deletable(client, resume_text):
    resume = ready_resume(client, resume_text)
    analysis = client.post('/api/ats/analyze', json={'resume_id': resume['id'], 'text': resume['text']}).json()
    with TestClient(app) as other:
        other.get('/api/config')
        assert other.get('/api/records').json()['total'] == 0
        assert other.get('/api/records/export?format=csv').status_code == 404
        assert other.delete('/api/records/' + analysis['id']).status_code == 404
    assert client.delete('/api/records/' + analysis['id']).status_code == 204
    assert client.get('/api/records').json()['total'] == 0
    assert client.get('/api/ats/results/' + analysis['id']).status_code == 404


def test_csv_export_neutralizes_formula_injection(client):
    resume = ready_resume(client, 'Alex Example\nalex@example.com\nExperience\nBuilt verified tools for internal reporting teams across regions.\nSkills\nPython, SQL', '=HYPERLINK(evil).txt')
    client.post('/api/ats/analyze', json={'resume_id': resume['id'], 'text': resume['text']})
    table = list(csv.DictReader(io.StringIO(client.get('/api/records/export?format=csv').content.decode('utf-8-sig'))))
    assert not table[0]['resume'].startswith('=')
