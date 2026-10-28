import asyncio
import io
import re
from collections import Counter
import pytest
import pymupdf
from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from app.main import app
from app.templates import TEMPLATES
from app.studio_api import sample_content
from app.document_renderer import render_pdf,render_document
from app.studio_docx import studio_docx_bytes
from app.editor_content import DraftContent,Section,DocumentStyle
from app.db import Session,SourceDocument
from app.config import settings


def words(value):return Counter(re.findall(r'\w+',value.lower()))


def test_eighteen_structurally_distinct_definitions():
 assert len(TEMPLATES)==18
 structures={(d['layout'],d['header'],d['heading'],d['entries'],d['skills'],tuple(d['rail']),tuple(d['lead'])) for d in TEMPLATES.values()}
 assert len(structures)==18
 assert all('ats_guidance' in d and 'docx_guidance' in d for d in TEMPLATES.values())


@pytest.mark.parametrize('key',list(TEMPLATES))
def test_every_template_real_pdf_and_supported_docx(key):
 content=sample_content();content.template=key
 pdf,info=asyncio.run(render_pdf(content))
 with pymupdf.open(stream=pdf,filetype='pdf') as doc:
  extracted=' '.join(p.get_text() for p in doc)
  assert len(doc)==info['pages']
  assert not (words(content.text())-words(extracted))
  assert any(p.get_links() for p in doc)
  assert abs(doc[0].rect.width-595.28)<1
 data=studio_docx_bytes(content);doc=Document(io.BytesIO(data));text=' '.join(node.text or '' for node in doc.element.body.iter(qn('w:t')))
 assert not (words(content.text())-words(text))
 if TEMPLATES[key]['columns']==2 or TEMPLATES[key]['layout']=='labelled':assert doc.tables
 assert doc.styles['Normal'].font.size.pt==TEMPLATES[key]['body_size']


def test_actual_multi_page_geometry_and_style_overrides():
 content=DraftContent(template='modern-professional',style=DocumentStyle(page_size='LETTER',body_size=12,heading_size=16,accent='#553344',margin_mm=18,page_limit=1),sections=[Section(id='contact',title='Contact information',html='<p>Actual Person</p>'),Section(id='experience',title='Experience',page_break_before=True,html=''.join('<p>Built verified tool '+str(i)+' for actual customer workflows and reliable internal reporting.</p>' for i in range(100)))])
 data,info=asyncio.run(render_pdf(content))
 assert info['pages']>2 and info['warnings']
 with pymupdf.open(stream=data,filetype='pdf') as pdf:
  assert all(abs(p.rect.width-612)<1 and abs(p.rect.height-792)<1 for p in pdf)
  assert 'tool 99' in ' '.join(p.get_text() for p in pdf)


def test_studio_creation_blank_change_template_preserves_content(client):
 created=client.post('/api/studio/create',json={'template':'two-column'})
 assert created.status_code==201,created.text
 value=created.json();assert 'Northstar' not in str(value) and 'Alex Morgan' not in str(value)
 assert client.get('/api/studio/'+value['resume_id']+'/export?format=pdf').status_code==422
 value['content']['sections'][0]['html']='<p>Actual Person</p><p>actual@example.com</p>'
 value['content']['sections'][1]['html']='<p>Built verified services for actual reporting workflows.</p>'
 saved=client.patch('/api/studio/'+value['resume_id'],json={'revision':value['revision'],'content':value['content']})
 assert saved.status_code==200,saved.text
 value=saved.json();source={s['id']:s['html'] for s in value['content']['sections']};value['content']['template']='sidebar-skills'
 saved=client.patch('/api/studio/'+value['resume_id'],json={'revision':value['revision'],'content':value['content']})
 assert saved.status_code==200
 assert {s['id']:s['html'] for s in saved.json()['content']['sections']}==source
 assert client.get('/api/studio/'+value['resume_id']+'/export?format=docx').status_code==200
 assert client.patch('/api/studio/'+value['resume_id'],json={'revision':1,'content':value['content']}).status_code==409


def test_retained_original_private_recreate_and_explicit_delete(client):
 source=Document();source.add_heading('Actual Person',0);source.add_paragraph('actual@example.com');source.add_heading('Experience',1);p=source.add_paragraph();p.add_run('Engineer | Actual Company').bold=True;source.add_paragraph('Built verified services for actual customer workflows.');buffer=io.BytesIO();source.save(buffer)
 response=client.post('/api/resumes/upload',files={'file':('actual.docx',buffer.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
 assert response.status_code==202
 rid=response.json()['id'];info=client.get('/api/imports/'+rid).json()
 assert info['original_available'] and info['profile'] and info['profile']['fonts'] is not None
 assert client.get('/api/imports/'+rid+'/original?download=true').content==buffer.getvalue()
 with TestClient(app) as other:
  other.get('/api/config');assert other.get('/api/imports/'+rid+'/original').status_code==404
 assert client.post('/api/imports/'+rid+'/recreate',json={'reviewed':False}).status_code==422
 recreated=client.post('/api/imports/'+rid+'/recreate',json={'reviewed':True,'template':'modern-professional'})
 assert recreated.status_code==201,recreated.text
 new_id=recreated.json()['resume_id'];assert new_id!=rid
 assert 'Actual Company' in client.get('/api/resumes/'+new_id).json()['text']
 assert client.delete('/api/imports/'+new_id+'/original').status_code==204
 assert client.get('/api/imports/'+new_id+'/original').status_code==410
 assert client.get('/api/editor/'+new_id).status_code==200
 assert client.delete('/api/resumes/'+rid).status_code==204
 with Session() as db:assert db.get(SourceDocument,rid) is None
 assert not (settings.storage_dir/'originals'/f'{rid}.docx').exists()


def test_unsafe_content_and_unsupported_style_rejected(client):
 result=client.post('/api/studio/preview',json={'content':{'template':'classic-ats','style':{'accent':'red; background:url(https://bad.invalid)'},'sections':[{'id':'a','title':'Summary','html':'<p>Actual</p>'}]}})
 assert result.status_code==422
 result=client.post('/api/studio/preview',json={'content':{'template':'classic-ats','style':{'columns':2},'sections':[{'id':'a','title':'Summary','html':'<p>Actual</p>'}]}})
 assert result.status_code==422
 content=DraftContent(sections=[Section(id='a',title='Summary',html='<p>Actual text</p><script>alert(1)</script><img src="https://bad.invalid">')])
 html=render_document(content)['html'];assert '<img src=' not in html and 'alert(1)' not in html
