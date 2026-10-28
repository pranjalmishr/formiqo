"""Short resumes must not spill onto extra pages (PDF canvas, Word export, PDF-to-Word).

Word layout itself needs LibreOffice/Word to measure, so these tests pin the causes:
explicit Word spacing instead of python-docx defaults, the canvas spill-over fit,
and the converter's source-derived page geometry.
"""
import asyncio
import io
import re
import pymupdf
import pytest
from docx import Document
from app.conversion import pdf_to_docx
from app.document_renderer import browser_path, render_pdf
from app.editor_content import DocumentStyle
from app.studio_api import sample_content
from app.studio_docx import studio_docx_bytes
from app.templates import TEMPLATES

chromium = pytest.mark.skipif(not browser_path(), reason='Chromium/Chrome is required for real canvas layout')


def test_word_export_overrides_python_docx_spacing_defaults():
    content = sample_content(); content.template = 'classic-ats'
    doc = Document(io.BytesIO(studio_docx_bytes(content)))
    heading, normal = doc.styles['Heading 1'], doc.styles['Normal']
    # python-docx defaults: 24 pt before Heading 1, 10 pt after Normal, 1.15 auto line height.
    assert heading.paragraph_format.space_before.pt <= 12
    assert normal.paragraph_format.space_after.pt <= 6
    assert normal.paragraph_format.line_spacing.pt <= TEMPLATES['classic-ats']['leading'] + .1


def test_word_column_tables_use_fixed_layout_without_padding():
    content = sample_content(); content.template = 'sidebar-skills'
    doc = Document(io.BytesIO(studio_docx_bytes(content)))
    xml = doc.tables[0]._tbl.xml
    assert 'w:tblLayout w:type="fixed"' in xml and 'w:tblCellMar' in xml


def test_word_export_applies_the_measured_fit():
    content = sample_content(); content.template = 'minimalist'
    normal = Document(io.BytesIO(studio_docx_bytes(content))).styles['Normal']
    fitted = Document(io.BytesIO(studio_docx_bytes(content, {'space': .4, 'leading': .92, 'body': .94, 'margin': .82}))).styles['Normal']
    assert fitted.paragraph_format.space_after.pt < normal.paragraph_format.space_after.pt
    assert fitted.font.size.pt >= 9


@chromium
def test_short_overflow_is_fitted_and_disclosed():
    content = sample_content(); content.template = 'classic-ats'
    data, info = asyncio.run(render_pdf(content))
    assert info['pages'] == 1 and info['fit'] and any('Tightened' in w for w in info['warnings'])
    with pymupdf.open(stream=data, filetype='pdf') as pdf:
        assert len(pdf) == 1
        sizes = [span['size'] for page in pdf for block in page.get_text('dict')['blocks'] for line in block.get('lines', []) for span in line['spans']]
        assert min(sizes) >= 7.9  # page numbers are 8 pt; body text never shrinks below 9 pt


@chromium
def test_fit_can_be_turned_off_and_long_documents_are_never_squeezed():
    content = sample_content(); content.template = 'classic-ats'; content.style = DocumentStyle(auto_fit=False)
    _, info = asyncio.run(render_pdf(content))
    assert info['pages'] == 2 and not info['fit']
    long = sample_content(); long.template = 'classic-ats'
    experience = next(s for s in long.sections if s.id == 'experience')
    experience.html = experience.html * 6
    _, info = asyncio.run(render_pdf(long))
    assert info['pages'] >= 2 and not info['fit']


@chromium
def test_pdf_to_word_uses_source_page_geometry(tmp_path):
    content = sample_content(); content.template = 'sidebar-skills'
    data, _ = asyncio.run(render_pdf(content))
    source = tmp_path / 'source.pdf'; source.write_bytes(data)
    output = tmp_path / 'converted.docx'
    warnings = pdf_to_docx(source, output)
    doc = Document(output)
    section = doc.sections[0]
    assert abs(section.page_width.pt - 595.28) < 1 and section.left_margin.pt < 72
    assert doc.styles['Normal'].paragraph_format.space_after.pt < 6
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert text.index('Technical Skills'.upper()) < text.index('Professional Summary'.upper())  # sidebar read first, not interleaved
    assert 'Python, TypeScript, React, PostgreSQL' in text and '1 / 1' not in text
    assert any(p.style.name == 'List Bullet' for p in doc.paragraphs)
    assert any('reading stream' in w for w in warnings)
