import io
import json
import zipfile
from datetime import timedelta
from pathlib import Path
import pytest
import pymupdf
from docx import Document
from PIL import Image
from fastapi.testclient import TestClient
from app.main import app
from app.db import Session,ConversionJob,DocumentRecord,now
from app.converter_api import cleanup_jobs,job_dir,execute_job
from app.conversion import matrix,page_indices,ConversionError,ConversionSettings,convert


def fixture_bytes(fmt):
    if fmt=='txt':return b'Actual first paragraph.\n\nActual second paragraph with verified text.'
    if fmt=='html':return b'<html><h2>Actual heading</h2><p>Actual paragraph.</p><script>alert(1)</script><img src="https://invalid.example/never-fetch"></html>'
    if fmt=='docx':
        doc=Document();doc.add_heading('Actual heading',level=1);doc.add_paragraph('Actual first paragraph.');doc.add_paragraph('Actual second paragraph.');buf=io.BytesIO();doc.save(buf);return buf.getvalue()
    if fmt=='pdf':
        doc=pymupdf.open();doc.new_page().insert_text((50,50),'PAGE ONE\nActual first paragraph.');doc.new_page().insert_text((50,50),'PAGE TWO\nActual second paragraph.');return doc.tobytes()
    img=Image.new('RGB',(160,80),'blue');buf=io.BytesIO();img.save(buf,format={'jpg':'JPEG','png':'PNG','webp':'WEBP'}[fmt]);return buf.getvalue()


def submit(client,source,target,data=None,options=None):
    from app.conversion import MIMES
    response=client.post('/api/converter/jobs',files=[('files',('input.'+source,fixture_bytes(source) if data is None else data,MIMES[source]))],data={'target':target,'options':json.dumps(options or {})})
    assert response.status_code==202,response.text
    job=client.get('/api/converter/jobs/'+response.json()['id']).json()
    return job


PAIRS=[(source,target) for source,targets in matrix().items() if source not in ('doc','pptx','xlsx') for target in targets]
@pytest.mark.parametrize('source,target',PAIRS)
def test_actual_conversion_outputs(client,source,target):
    job=submit(client,source,target)
    assert job['status']=='completed',job
    assert job['outputs'] and job['duration_ms'] is not None
    for index,out in enumerate(job['outputs']):
        result=client.get('/api/converter/jobs/'+job['id']+'/download?index='+str(index))
        assert result.status_code==200 and len(result.content)==out['size']
        if out['format']=='pdf':
            with pymupdf.open(stream=result.content,filetype='pdf') as doc: assert len(doc)>0
        elif out['format']=='docx':
            doc=Document(io.BytesIO(result.content));assert any(p.text.strip() for p in doc.paragraphs)
        elif out['format'] in ('jpg','png','webp'):
            with Image.open(io.BytesIO(result.content)) as img: img.verify()
        elif out['format']=='txt': assert 'Actual' in result.text and '<script>' not in result.text
        elif out['format']=='html': assert '<script>' not in result.text and 'Actual' in result.text
        elif out['format']=='zip':
            with zipfile.ZipFile(io.BytesIO(result.content)) as z:assert len(z.namelist())==2 and z.testzip() is None
    assert not (job_dir(job['id'])/'inputs').exists()
    assert not (job_dir(job['id'])/'request.json').exists()
    items=client.get('/api/documents?category=converted').json()['items']
    assert any(item['job_id']==job['id'] for item in items)


def test_pdf_to_word_page_order(client):
    job=submit(client,'pdf','docx')
    assert job['status']=='completed',job
    data=client.get('/api/converter/jobs/'+job['id']+'/download').content
    text='\n'.join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    assert text.index('PAGE ONE')<text.index('PAGE TWO')
    assert any('not an exact' in warning for warning in job['warnings'])


def test_image_order_and_aspect_ratio(client):
    from app.conversion import MIMES
    red=Image.new('RGB',(160,80),'red');red_data=io.BytesIO();red.save(red_data,format='PNG')
    files=[('files',('first.png',red_data.getvalue(),MIMES['png'])),('files',('second.jpg',fixture_bytes('jpg'),MIMES['jpg']))]
    response=client.post('/api/converter/jobs',files=files,data={'target':'pdf','options':json.dumps({'orientation':'landscape','page_size':'LETTER','margin_mm':10})})
    assert response.status_code==202,response.text
    job=client.get('/api/converter/jobs/'+response.json()['id']).json();assert job['status']=='completed',job
    result=client.get('/api/converter/jobs/'+job['id']+'/download')
    with pymupdf.open(stream=result.content,filetype='pdf') as doc:
        assert len(doc)==2 and doc[0].rect.width>doc[0].rect.height
        images=doc[0].get_image_info();box=images[0]['bbox'];assert abs((box[2]-box[0])/(box[3]-box[1])-2)<.01
        a=doc[0].get_pixmap();b=doc[1].get_pixmap();first=a.pixel(a.width//2,a.height//2);second=b.pixel(b.width//2,b.height//2)
        assert first[0]>240 and first[2]<15 and second[2]>240 and second[0]<15


def test_page_selection_and_invalid_range(client):
    job=submit(client,'pdf','png',options={'pages':'2','dpi':72})
    assert [o['filename'] for o in job['outputs']]==['page-002.png','all-pages.zip']
    invalid=submit(client,'pdf','png',options={'pages':'4'})
    assert invalid['status']=='failed' and invalid['error_code']=='invalid_page_range'
    assert client.get('/api/converter/jobs/'+invalid['id']+'/download').status_code==409
    with pytest.raises(ConversionError):page_indices('1; rm',2)


def test_corrupt_empty_unsupported_and_encrypted(client):
    response=client.post('/api/converter/jobs',files={'files':('x.pdf',b'not pdf','application/pdf')},data={'target':'txt'})
    assert response.status_code==422
    response=client.post('/api/converter/jobs',files={'files':('x.txt',b'valid','text/plain')},data={'target':'webp'})
    assert response.status_code==415
    response=client.post('/api/converter/jobs',files={'files':('x.txt',b'','text/plain')},data={'target':'pdf'})
    assert response.status_code==422
    job=submit(client,'pdf','txt',data=b'%PDF-1.7\ninvalid')
    assert job['status']=='failed'
    doc=pymupdf.open();doc.new_page().insert_text((50,50),'Secret text')
    encrypted=doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256,owner_pw='secret',user_pw='secret')
    job=submit(client,'pdf','docx',data=encrypted)
    assert job['status']=='failed' and job['error_code']=='password_protected'


def test_private_download_expiry_delete_and_rename(client):
    job=submit(client,'txt','pdf')
    with TestClient(app) as other:
        other.get('/api/config')
        assert other.get('/api/converter/jobs/'+job['id']).status_code==404
        assert other.get('/api/converter/jobs/'+job['id']+'/download').status_code==404
        assert other.delete('/api/converter/jobs/'+job['id']).status_code==404
    document=client.get('/api/documents?category=converted').json()['items'][0]
    assert client.patch('/api/documents/'+document['id'],json={'filename':'Renamed output.pdf'}).status_code==200
    assert client.patch('/api/documents/'+document['id'],json={'filename':'../../escape.pdf'}).status_code==422
    with Session() as db:
        row=db.get(ConversionJob,job['id']);row.expires_at=now()-timedelta(minutes=1);db.commit()
    assert client.get('/api/converter/jobs/'+job['id']+'/download').status_code==404
    cleanup_jobs();assert not job_dir(job['id']).exists()
    assert client.get('/api/documents/'+document['id']).status_code==404
    job=submit(client,'txt','docx');directory=job_dir(job['id'])
    assert client.delete('/api/converter/jobs/'+job['id']).status_code==204
    assert not directory.exists()


def test_timeout_and_missing_capabilities(client,tmp_path,monkeypatch):
    class FakeProcess:
        pid=123;returncode=0
        def communicate(self,timeout=None):
            if timeout:raise __import__('subprocess').TimeoutExpired('worker',timeout)
            return b'',b''
    monkeypatch.setattr('app.converter_api.subprocess.Popen',lambda *a,**k:FakeProcess())
    killed=[];monkeypatch.setattr('app.converter_api.os.killpg',lambda *args:killed.append(args))
    result=execute_job('00000000-0000-0000-0000-000000000000')
    assert result['error_code']=='timeout' and killed
    supported=client.get('/api/converter/supported-formats').json()
    assert 'doc' not in supported['matrix']
    assert not supported['office_available']


def test_scanned_pdf_ocr_path_and_actionable_missing_ocr(tmp_path,monkeypatch):
    img=Image.new('RGB',(400,300),'white');buf=io.BytesIO();img.save(buf,format='PNG')
    doc=pymupdf.open();doc.new_page().insert_image(pymupdf.Rect(0,0,400,300),stream=buf.getvalue());source=tmp_path/'scan.pdf';doc.save(source)
    work=tmp_path/'out';work.mkdir()
    monkeypatch.setattr('app.conversion.ocr_page',lambda page:'Actual OCR text from synthetic scan')
    result=convert([source],'pdf','docx',work,ConversionSettings())
    assert any('OCR used' in warning for warning in result['warnings'])
    assert 'Actual OCR' in '\n'.join(p.text for p in Document(work/'converted.docx').paragraphs)
    monkeypatch.setattr('app.conversion.ocr_page',lambda page:(_ for _ in ()).throw(__import__('app.extraction',fromlist=['ExtractionError']).ExtractionError('Tesseract missing')))
    with pytest.raises(ConversionError,match='Tesseract'):convert([source],'pdf','docx',work,ConversionSettings())
