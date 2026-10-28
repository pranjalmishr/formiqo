import asyncio
import io
from datetime import timedelta
from pathlib import Path
import json
import zipfile
import pymupdf as fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient
from app.main import app, expire_records, run_worker
from app.db import Session, Resume, now
from app.config import settings
from app.scoring import analyze, keywords
from app.extraction import extract, ExtractionError
from app.schemas import SemanticResponse
from app import ai


def upload(client, data, name='resume.txt', mime='text/plain'):
    response = client.post('/api/resumes/upload', files={'file': (name, data, mime)})
    assert response.status_code == 202, response.text
    identifier = response.json()['id']
    return client.get('/api/resumes/' + identifier).json()


def analyze_api(client, resume, text=None, **extra):
    return client.post('/api/ats/analyze', json={'resume_id': resume['id'], 'text': text or resume['text'], **extra})


def test_txt_upload_analysis_edit_compare(client, resume_text):
    r = upload(client, resume_text.encode())
    assert r['status'] == 'ready'
    assert 'Python' in r['text']
    result = analyze_api(client, r).json()
    assert result['overall_score'] == sum(c['score'] for c in result['categories'])
    assert sum(c['maximum'] for c in result['categories']) == 100
    assert all(0 <= c['score'] <= c['maximum'] for c in result['categories'])
    assert not result['semantic_observations']
    original = r['original_text']
    edited = r['text'].replace('jordan@example.com', '')
    saved = client.patch('/api/resumes/' + r['id'], json={'text': edited})
    assert saved.status_code == 200
    assert saved.json()['original_text'] == original
    updated = analyze_api(client, r, edited).json()
    comparison = client.post('/api/ats/compare', json={'original_id': result['id'], 'updated_id': updated['id']}).json()
    assert comparison['change'] == updated['overall_score'] - result['overall_score']
    assert comparison['change'] < 0
    assert len(client.get('/api/resumes/' + r['id'] + '/versions').json()['items']) >= 2
    assert client.get('/api/resumes/' + r['id'] + '/download').text == edited
    assert not list(settings.storage_dir.glob(r['id'] + '.*'))


def test_pdf_upload_preserves_page_order(client, resume_text):
    doc = fitz.open()
    doc.new_page().insert_text((50, 50), 'PAGE ONE\n' + resume_text, fontsize=9)
    doc.new_page().insert_text((50, 50), 'PAGE TWO\nCertifications\nVerified training', fontsize=10)
    r = upload(client, doc.tobytes(), 'resume.pdf', 'application/pdf')
    assert r['status'] == 'ready'
    assert r['text'].index('PAGE ONE') < r['text'].index('PAGE TWO')
    assert r['extraction']['page_count'] == 2


def test_docx_upload_table_order(client, resume_text):
    doc = Document()
    doc.add_paragraph('START')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'TABLE CONTENT'
    table.cell(0, 1).text = 'RIGHT CELL'
    doc.add_paragraph('END')
    doc.add_paragraph(resume_text)
    out = io.BytesIO()
    doc.save(out)
    r = upload(client, out.getvalue(), 'resume.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    assert r['status'] == 'ready'
    assert r['text'].index('START') < r['text'].index('TABLE CONTENT') < r['text'].index('END')
    assert any(w['code'] == 'tables' for w in r['extraction']['warnings'])


@pytest.mark.parametrize('name,mime,data,status', [
    ('a.pdf','application/pdf',b'not a pdf',422),
    ('a.docx','application/vnd.openxmlformats-officedocument.wordprocessingml.document',b'PKbroken',422),
    ('a.doc','application/msword',b'data',415),
    ('a.txt','text/plain',b'',422),
    ('a.txt','text/plain',b'\x00binary',422),
    ('a.pdf','text/plain',b'%PDF-data',415),
])
def test_upload_rejections(client, name, mime, data, status):
    r = client.post('/api/resumes/upload', files={'file':(name,data,mime)})
    assert r.status_code == status


def test_corrupted_and_empty_documents(client):
    r = upload(client, b'%PDF-1.7\ninvalid', 'broken.pdf', 'application/pdf')
    assert r['status'] == 'failed' and 'could not be read' in r['error']
    r = upload(client, b'   \n\n ', 'empty.txt')
    assert r['status'] == 'failed' and 'No readable text' in r['error']


def test_password_protected_pdf(client):
    doc = fitz.open()
    doc.new_page().insert_text((50,50),'Private')
    encrypted = doc.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw='secret', user_pw='secret')
    r = upload(client, encrypted, 'private.pdf', 'application/pdf')
    assert r['status'] == 'failed' and 'Password-protected' in r['error']


def test_ownership_download_and_deletion(client, resume_text):
    r = upload(client, resume_text.encode())
    result = analyze_api(client, r).json()
    with TestClient(app) as other:
        other.get('/api/config')
        for path in ['/api/resumes/'+r['id'], '/api/resumes/'+r['id']+'/download', '/api/ats/results/'+result['id']]:
            assert other.get(path).status_code == 404
        assert other.delete('/api/resumes/'+r['id']).status_code == 404
    assert client.delete('/api/resumes/'+r['id']).status_code == 204
    assert client.get('/api/ats/results/'+result['id']).status_code == 404
    assert client.get('/api/resumes/'+r['id']).status_code == 404


def test_expiration(client, resume_text):
    r = upload(client, resume_text.encode())
    result = analyze_api(client, r).json()
    with Session() as db:
        row = db.get(Resume, r['id'])
        row.expires_at = now() - timedelta(minutes=1)
        db.commit()
    assert client.get('/api/resumes/'+r['id']+'/download').status_code == 404
    expire_records()
    assert client.get('/api/ats/results/'+result['id']).status_code == 404
    with Session() as db:
        assert db.get(Resume, r['id']) is None


def test_scoring_determinism_and_no_optional_penalties(resume_text):
    first = analyze(resume_text, {})
    assert first == analyze(resume_text, {})
    student = analyze('Alex Student\nalex@example.com\nSummary\nLearning useful engineering techniques and exploring tools.\nEducation\nCurrent student working on coursework and practical assignments.', {})
    assert next(c for c in student['categories'] if c['key']=='experience')['score'] == 20
    assert next(c for c in student['categories'] if c['key']=='skills')['score'] == 15
    assert not any('degree' in f['title'].lower() for f in student['findings'])


def test_jd_missing_terms_and_required_preferred(resume_text):
    jd = 'Required: Python and AWS.\nPreferred: Kubernetes.\nResponsibilities: Build services using Python and AWS.'
    k = keywords(jd, resume_text)
    assert 'python' in k['found'] and 'aws' in k['missing'] and 'kubernetes' in k['missing']
    assert k['groups']['required'] and k['groups']['preferred']
    result = analyze(resume_text, {}, jd, 'job')
    assert result['mode'] == 'job'
    assert any(f['requires_confirmation'] for f in result['findings'] if f['category']=='keywords')


def test_job_validation_and_incompatible_compare(client, resume_text):
    r = upload(client, resume_text.encode())
    assert analyze_api(client, r, mode='job', job_description='tiny').status_code == 422
    a = analyze_api(client,r).json()
    b = analyze_api(client,r,mode='job',job_description='Required: Python engineering experience. Preferred: SQL and AWS infrastructure.').json()
    assert client.post('/api/ats/compare',json={'original_id':a['id'],'updated_id':b['id']}).status_code==409


def test_scanned_detection_and_ocr_fallback(tmp_path, monkeypatch):
    from PIL import Image, ImageDraw
    image = Image.new('RGB',(800,500),'white')
    ImageDraw.Draw(image).text((30,30),'SCANNED RESUME',fill='black')
    stream = io.BytesIO()
    image.save(stream,format='PNG')
    pdf = fitz.open()
    pdf.new_page().insert_image(fitz.Rect(0,0,600,400),stream=stream.getvalue())
    path = tmp_path/'scan.pdf'
    pdf.save(path)
    monkeypatch.setattr('app.extraction.ocr_page',lambda page:'Scanned Resume\nuser@example.com\nEducation\nCurrent student')
    result = extract(path,'pdf')
    assert result['extraction']['ocr_pages'] == [1]
    assert 'Current student' in result['text']
    monkeypatch.setattr('app.extraction.shutil.which',lambda _:None)
    monkeypatch.undo()
    monkeypatch.setattr('app.extraction.shutil.which',lambda _:None)
    with pytest.raises(ExtractionError,match='Tesseract'):
        extract(path,'pdf')


def test_schema_rejects_ai_numbers():
    with pytest.raises(ValueError):
        SemanticResponse.model_validate({'observations':[], 'score':99})
    with pytest.raises(ValueError):
        SemanticResponse.model_validate({'observations':[{'category':'salary','evidence':'invented','recommendation':'x','strength':True}]})


def test_ai_evidence_filter(monkeypatch, resume_text):
    monkeypatch.setattr(settings,'llm_api_key','test-key')
    class FakeResponse:
        def raise_for_status(self): pass
        def json(self): return {'choices':[{'message':{'content':json.dumps({'observations':[{'category':'experience','evidence':'Developed Python services','recommendation':'Explain the actual outcome without inventing metrics.','strength':True},{'category':'experience','evidence':'Made a million dollars','recommendation':'This is not in the document.','strength':True}]})}}]}
    class FakeClient:
        def __init__(self,**kwargs): pass
        async def __aenter__(self): return self
        async def __aexit__(self,*args): pass
        async def post(self,*args,**kwargs): return FakeResponse()
    monkeypatch.setattr(ai.httpx,'AsyncClient',FakeClient)
    observations,notice = asyncio.run(ai.evaluate(resume_text,''))
    assert len(observations)==1 and observations[0]['evidence'] in resume_text


def test_origin_and_size_limits(client, monkeypatch):
    assert client.post('/api/resumes/upload',headers={'Origin':'https://evil.example'},files={'file':('a.txt',b'valid','text/plain')}).status_code==403
    monkeypatch.setattr(settings,'max_upload_mb',1)
    assert client.post('/api/resumes/upload',files={'file':('a.txt',b'a'*(1024*1024+1),'text/plain')}).status_code==413


def test_worker_timeout(tmp_path, monkeypatch):
    class Process:
        pid=123
        def communicate(self,timeout=None):
            if timeout: raise __import__('subprocess').TimeoutExpired('worker',timeout)
            return b'',b''
    monkeypatch.setattr('app.main.subprocess.Popen',lambda *a,**k:Process())
    killed=[]
    monkeypatch.setattr('app.main.os.killpg',lambda *args:killed.append(args))
    result=run_worker(tmp_path/'fake.pdf','pdf')
    assert 'timed out' in result['error'] and killed


def test_external_ai_requires_explicit_consent(client, resume_text, monkeypatch):
    calls = []
    async def evaluator(text, jd):
        calls.append(text)
        return [], 'AI schema validated'
    monkeypatch.setattr('app.main.evaluate', evaluator)
    r = upload(client, resume_text.encode())
    assert analyze_api(client, r).status_code == 201
    assert calls == []
    assert analyze_api(client, r, ai_consent=True).status_code == 201
    assert calls == [r['text']]


def test_request_body_limit(client):
    response = client.post('/api/ats/analyze', content=b'x' * (1024 * 1024 + 1), headers={'Content-Type': 'application/json'})
    assert response.status_code == 413


def test_configuration_rejects_invalid_weights():
    from app.config import Settings
    with pytest.raises(ValueError):
        Settings(scoring_weights={'structure': 100})
    with pytest.raises(ValueError):
        Settings(retention_hours=0)
