"""Native Firestore regression tests. Emulator integration never targets a real project.

Run: FIRESTORE_EMULATOR_HOST=127.0.0.1:8085 python -m pytest tests/test_firestore.py
The ordinary SQL regression suite is kept unchanged and runs alongside these checks.
"""
from datetime import timedelta
import io
import os
import uuid

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from sqlalchemy.exc import IntegrityError

from app.config import settings
from app.db import Resume,User,AuthToken,Session,now,uid
from app.firestore_store import FirestoreUnitOfWork,client_for,encode_record,decode_record
from app.main import app
from app.persistence import find,consume_token_id


def test_small_records_remain_normal_firestore_fields():
    value={'owner':'synthetic','score':83,'text':'Actual synthetic resume','at':now()}
    assert encode_record(value)==value and decode_record(value)==value


def test_large_unicode_records_round_trip_without_truncation():
    value={'owner':'synthetic','id':uid(),'text':'\u03a9\U0001f680'*200000,'at':now()}
    packed=encode_record(value)
    assert '_formiqo_payload_zlib' in packed and packed['owner']=='synthetic'
    assert decode_record(packed)==value


def test_document_size_limit_is_explicit():
    with pytest.raises(HTTPException) as error:encode_record({'text':'x'*4_100_000})
    assert error.value.status_code==413


@pytest.fixture
def firestore_app(monkeypatch):
    if not os.environ.get('FIRESTORE_EMULATOR_HOST'):
        pytest.skip('Firestore emulator is not running')
    monkeypatch.setattr(settings,'database_backend','firestore')
    monkeypatch.setattr(settings,'firebase_project_id','demo-formiqo')
    monkeypatch.setattr(settings,'firebase_credentials_file',None)  # emulator never refreshes a real OAuth credential
    monkeypatch.setattr(settings,'firebase_namespace','test_'+uuid.uuid4().hex)
    from app.storage import _rates,_lock
    with _lock:_rates.clear()
    client=client_for('demo-formiqo')
    with TestClient(app) as browser:
        browser.get('/api/config')
        yield browser
    assert client.project=='demo-formiqo'
    client.recursive_delete(client.collection('formiqo').document(settings.firebase_namespace))


def create(browser):
    response=browser.post('/api/builder/create',json={'contact':{'name':'Firestore Example','email':'firestore@example.com'},'summary':'Developed verified Python services for actual internal reporting workflows.','skills':'Python, SQL, React','experience':[{'title':'Engineer','employer':'Synthetic Company','highlights':'Built verified Python services for 12 internal teams.'}]})
    assert response.status_code==201,response.text
    return response.json()


def test_firestore_is_primary_storage_for_builder_editor_and_records(firestore_app):
    browser=firestore_app;draft=create(browser);rid=draft['resume_id']
    assert browser.get('/api/resumes/'+rid).json()['text']
    payload={'revision':draft['content'].get('revision',1),'content':draft['content'],'change_summary':'Native Firestore edit'}
    payload['content']['sections'][1]['html']='<p>Improved verified Python services for 15 internal reporting teams.</p>'
    saved=browser.patch('/api/editor/'+rid,json=payload)
    assert saved.status_code==200,saved.text
    assert saved.json()['revision']==2
    # The stale revision is rejected instead of overwriting the Firestore record.
    payload['content']['sections'][1]['html']='<p>Stale content must not win.</p>'
    assert browser.patch('/api/editor/'+rid,json=payload).status_code==409
    text=browser.get('/api/resumes/'+rid).json()['text']
    assert '15 internal' in text and 'Stale content' not in text
    general=browser.post('/api/ats/analyze',json={'resume_id':rid,'text':text})
    assert general.status_code==201,general.text
    matched=browser.post('/api/ai/match-job-description',json={'resume_id':rid,'job_title':'Engineer','company':'Synthetic Co','job_description':'Requirements: Python, SQL and Terraform. Responsibilities: build reliable services and reports. Preferred: React.'})
    assert matched.status_code==201,matched.text
    records=browser.get('/api/records').json()
    assert records['total']==2 and records['summary']['job_matches']==1
    response=browser.get('/api/records/export?format=xlsx')
    from openpyxl import load_workbook
    assert load_workbook(io.BytesIO(response.content))['Score records'].max_row==3
    assert browser.get('/api/activity').json()['items']
    with Session() as store:
        loaded=store.get(Resume,rid)
        assert loaded is not None and store.is_firestore
        assert store.reference(Resume,rid).get().exists
    assert browser.delete('/api/resumes/'+rid).status_code==204
    assert browser.get('/api/records').json()['total']==0


def test_firestore_ownership_and_guest_claim_survive_sign_in(firestore_app):
    browser=firestore_app;draft=create(browser);rid=draft['resume_id']
    with TestClient(app) as other:
        other.get('/api/config')
        assert other.get('/api/resumes/'+rid).status_code==404
        assert other.get('/api/records').json()['total']==0
    data={'name':'Firestore Example','email':'free-plan@example.com','password':'verified synthetic password 2026!','claim_guest':True}
    signed=browser.post('/api/auth/register',json=data)
    assert signed.status_code==201,signed.text
    assert browser.get('/api/resumes/'+rid).status_code==200
    assert browser.post('/api/auth/logout').status_code==204
    assert browser.get('/api/resumes/'+rid).status_code==404
    login=browser.post('/api/auth/login',json={'email':data['email'],'password':data['password']})
    assert login.status_code==200,login.text
    assert browser.get('/api/resumes/'+rid).status_code==200
    assert browser.request('DELETE','/api/auth/account',json={'password':data['password']}).status_code==204
    assert browser.get('/api/resumes/'+rid).status_code==404


def test_firestore_transactions_prevent_lost_updates(firestore_app):
    draft=create(firestore_app);rid=draft['resume_id']
    with Session() as left,Session() as right:
        first=left.get(Resume,rid);second=right.get(Resume,rid)
        first.filename='First transaction wins';left.commit()
        second.filename='Second must conflict'
        with pytest.raises(HTTPException) as error:right.commit()
        assert error.value.status_code==409
    assert firestore_app.get('/api/resumes/'+rid).json()['filename']=='First transaction wins'


def test_firestore_unique_emails_are_transactional(firestore_app):
    users=[User(name='Synthetic',email='unique@example.com',password_hash='not-a-real-password-hash') for _ in range(2)]
    with Session() as left,Session() as right:
        left.add(users[0]);right.add(users[1]);left.flush();right.flush()
        left.commit()
        with pytest.raises(IntegrityError):right.commit()
    with Session() as store:
        assert len(find(store,User,[('email','==','unique@example.com')]))==1


def test_firestore_one_time_token_consumption_is_atomic(firestore_app):
    with Session() as store:
        user=User(name='Token Example',email='token@example.com',password_hash='synthetic-only')
        store.add(user);store.add(AuthToken(token_hash='a'*64,user_id=user.id,purpose='reset',expires_at=now()+timedelta(minutes=30)));store.commit()
    with Session() as left,Session() as right:
        assert consume_token_id(left,AuthToken,'a'*64,'reset',now())==user.id
        assert consume_token_id(right,AuthToken,'a'*64,'reset',now())==user.id
        left.commit()
        with pytest.raises(HTTPException) as error:right.commit()
        assert error.value.status_code==409
    with Session() as store:assert consume_token_id(store,AuthToken,'a'*64,'reset',now()) is None


def test_firestore_real_converter_metadata_and_download(firestore_app):
    browser=firestore_app
    response=browser.post('/api/converter/jobs',files={'files':('synthetic.txt',b'Actual synthetic document text for a conversion.','text/plain')},data={'target':'docx'})
    assert response.status_code==202,response.text
    job=browser.get('/api/converter/jobs/'+response.json()['id']).json()
    assert job['status']=='completed',job
    result=browser.get('/api/converter/jobs/'+job['id']+'/download')
    assert result.status_code==200 and result.content[:2]==b'PK'
    assert browser.get('/api/documents?category=converted').json()['items']
    assert browser.delete('/api/converter/jobs/'+job['id']).status_code==204


def test_firestore_rejects_orphaned_child_records(firestore_app):
    from app.db import ResumeVersion
    with Session() as store:
        store.add(ResumeVersion(resume_id=uid(),text='No parent exists'))
        with pytest.raises(IntegrityError):store.commit()


def test_firestore_migration_is_dry_run_by_default_idempotent_and_non_destructive(firestore_app,tmp_path):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session as SQLSession
    from app.db import Base,Analysis,ResumeVersion
    from app.migrate_firestore import migrate
    url='sqlite:///'+str(tmp_path/'source.db');engine=create_engine(url)
    Base.metadata.create_all(engine)
    rid=uid();expired=uid();analysis_id=uid()
    with SQLSession(engine) as source:
        source.add_all([Resume(id=rid,owner='migration-owner',filename='synthetic.txt',original_format='txt',text='Actual synthetic text',expires_at=now()+timedelta(hours=2)),Resume(id=expired,owner='migration-owner',filename='expired.txt',original_format='txt',expires_at=now()-timedelta(days=2)),Analysis(id=analysis_id,resume_id=rid,owner='migration-owner',overall_score=72,result={'mode':'general'},scoring_version='synthetic-test'),ResumeVersion(resume_id=expired,text='Expired child')]);source.commit()
    client=client_for('demo-formiqo');namespace=settings.firebase_namespace
    result=migrate('demo-formiqo',url,namespace,client=client)
    assert result['new_records']==2 and not result['applied']
    with Session() as store:assert store.get(Resume,rid) is None
    result=migrate('demo-formiqo',url,namespace,apply=True,client=client)
    assert result['applied'] and result['verified_records']==2 and result['local_backup_created']
    assert migrate('demo-formiqo',url,namespace,apply=True,client=client)['new_records']==0
    with SQLSession(engine) as source:assert source.get(Resume,rid).text=='Actual synthetic text'
    with Session() as store:
        row=store.get(Resume,rid);row.filename='Newer target content';store.commit()
    with pytest.raises(RuntimeError,match='differs from the source'):migrate('demo-formiqo',url,namespace,apply=True,client=client)
    engine.dispose()
